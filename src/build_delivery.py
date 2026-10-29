"""Build a two-page report, five-slide presentation and a curated submission ZIP."""
from pathlib import Path
import hashlib
import json
import subprocess
import zipfile
import os
import shutil

os.environ.setdefault('MPLCONFIGDIR','/tmp/pvc-matplotlib')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4

ROOT=Path(__file__).resolve().parents[1]
ART=ROOT/'artifacts'
OUT=ROOT/'deliverables'


def narrative(meta):
    delta=meta['final_naive_recent8_smape']-meta['final_model_recent8_smape']
    strategy='Все 13 позиций равны нулю; P&L и просадка равны нулю. Это отсутствие сделок, а не прибыльная торговая система.' if meta['final_strategy']=='flat' else (
        f"Выбрана стратегия с фильтром ошибки и ограничением: максимум одна активная неделя на 13-недельный блок. "
        f"Исторический P&L {meta['strategy_pnl']:+.5f}; просадка {meta['strategy_drawdown']:.5f}. "
        f"С условными затратами 0,1% на оборот: {meta['strategy_stress_pnl']:+.5f}. "
        'В итоговом файле одна короткая позиция на первой неделе, остальные нулевые.')
    return dict(gap=delta,strategy=strategy,
                recent=f"Последние 8 окон: {meta['final_model_recent8_smape']:.2f}% sMAPE против {meta['final_naive_recent8_smape']:.2f}% у naive; отрыв {delta:+.3f} п.п.",
                broad=f"На 20 блоках фиксированная модель: {meta['final_model_mean_smape']:.2f}% против {meta['final_naive_mean_smape']:.2f}% у naive. Устойчивое превосходство не подтверждено.")


def build_pdf(meta):
    font=Path('/usr/share/fonts/truetype/dejavu')
    pdfmetrics.registerFont(TTFont('DejaVu',str(font/'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('DejaVu-Bold',str(font/'DejaVuSans-Bold.ttf')))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='BodyRU',fontName='DejaVu',fontSize=9,leading=13,spaceAfter=7))
    styles.add(ParagraphStyle(name='TitleRU',fontName='DejaVu-Bold',fontSize=17,leading=21,spaceAfter=13,textColor=colors.HexColor('#173c66')))
    styles.add(ParagraphStyle(name='HeadingRU',fontName='DejaVu-Bold',fontSize=11,leading=15,spaceBefore=8,spaceAfter=5))
    styles.add(ParagraphStyle(name='CellRU',fontName='DejaVu',fontSize=8,leading=11))
    story=[]
    text=narrative(meta)
    def p(s):story.append(Paragraph(s,styles['BodyRU']))
    def h(s):story.append(Paragraph(s,styles['HeadingRU']))
    story.append(Paragraph('Прогноз цены ПВХ и поддержка решений',styles['TitleRU']))
    p('Команда Roshim_hakaton · финальная версия 29 октября 2026 года')
    h('Задача и данные')
    p('Цель — индекс k недельного ряда «ПВХ SG-5 (Китай) DAP Достык, USD/т». Прогноз фиксируется на 13 недель: 10.07–02.10.2026. Используется вся предоставленная обучающая история; будущие факты неизвестны.')
    p('Официальный XLSB: 36 293 записи, 170 рядов. Цель: 552 исходные записи, 551 наблюдаемая неделя после объединения повторного релиза и 13 пропусков на сетке. Неоднозначные контрактные котировки исключены; давность внешних значений ограничена 45 днями.')
    h('Метод и проверка')
    p('Сравнены 16 исходных вариантов: naive, Holt, Ridge, CatBoost и ансамбль. Финальная модель — CatBoost с абсолютной ошибкой, лагами, волатильностью, сезонностью и рыночными признаками. Отбор признаков и обучение повторяются внутри исторического префикса. Будущие цены сырья не используются.')
    p('24 неперекрывающихся блока по 13 недель. Первые 4 используются для разогрева исследования последовательного выбора; 20 последующих — для сравнения. Последние 8 окон использованы для выбора кандидата. Результаты относятся к разработческой исторической проверке, не к независимому конкурсному тесту.')
    h('Результаты точности')
    p(text['recent']);p(text['broad'])
    comparison=pd.read_csv(ART/'final_model_comparison.csv')
    data=[[Paragraph(x,styles['CellRU']) for x in ['Кандидат','Средний sMAPE','Худший блок','Последние 8']]]
    labels={'catboost_mae_external':'Финальный CatBoost','naive':'Naive','conservative_25':'75% naive + 25% MAE','conservative_50':'50% naive + 50% MAE'}
    for r in comparison.itertuples():
        data.append([Paragraph(labels[r.model],styles['CellRU']),f'{r.mean:.3f}%',f'{r.worst:.3f}%',f'{r.recent8:.3f}%'])
    table=Table(data,colWidths=[205,90,90,90])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7eef7')),('FONTNAME',(0,0),(-1,-1),'DejaVu'),('FONTSIZE',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#173c66'))]))
    story.append(table)
    p('Два осторожных ансамбля проверены перед сдачей. Они улучшают среднюю и худшую ошибку, но ухудшают недавние окна. По заранее записанному правилу без ухудшения проверяемых показателей замена отклонена, исходный прогноз сохранён.')
    story.append(PageBreak())
    story.append(Paragraph('Риск, применение и воспроизводимость',styles['TitleRU']))
    h('Интервалы и горизонт')
    p(f"Номинальные 90%-е эмпирические интервалы используют только ошибки завершённых блоков и текущую известную волатильность. Для фиксированного финального кандидата покрытие составляет {meta['interval_coverage']:.1%} на {meta['interval_n']} наблюдениях. Высокое покрытие не равно точности: интервалы широкие, будущее покрытие не гарантировано.")
    p('Ошибка исследована отдельно по каждому горизонту на одинаковых полных блоках. В дополнительном парном исследовании последовательного выбора уточнение t+2 после факта t+1 снизило sMAPE с 4,44 до 2,87% на 15 случаях. Это сценарий эксплуатации с новыми данными, не конкурсный фиксированный прогноз.')
    h('Стратегия')
    p(text['strategy'])
    p('Сигнал — прогнозное недельное изменение, превышающее максимум из 0,5% последней цены и половины медианной исторической ошибки изменения для этого горизонта. Ошибки берутся только из завершённых блоков. Из допущенных сигналов выбирается один с максимальным отношением изменения к порогу; будущий P&L не используется.')
    p(f"На 20 блоках: {meta.get('strategy_active_weeks',0)} активных недель, {meta.get('strategy_profitable_blocks',0)} прибыльных и {meta.get('strategy_losing_blocks',0)} убыточных блоков. Доля лучшего блока в положительных результатах {meta.get('strategy_largest_positive_fraction',0):.1%}. С затратами после исключения лучшего блока остаётся {meta.get('strategy_stress_without_best_block',0):+.5f}. Это проверка на уже исследованной истории, не независимое доказательство будущей прибыли.")
    p('Исторический buy & hold: −0,82443 условных единиц. Индекс k не переводится в рубли; объёмы и официальные комиссии не заданы. Условные затраты 0,1% на оборот — исследовательская чувствительность, не правило жюри. Число сделок невелико; будущая прибыль не гарантируется.')
    h('Бизнес-применение и ограничения')
    p('Инструмент показывает ценовой ориентир и диапазон неопределённости для обсуждения сроков продаж, закупок и планирования. DAP Достык не заменяет конкретную контрактную цену. Модельная важность признака не доказывает причинность. Основные ограничения — смена рыночных режимов, небольшое число временных блоков и отсутствие точного времени публикации котировок.')
    h('Состав и запуск')
    p('В комплекте: forecast.csv с forecast_lower/forecast_upper, strategy.csv, исходный код, README, модель CatBoost, этот отчёт и презентация. CSV: разделитель ;, UTF-8 BOM, официальные пятницы. Прогноз и позиции фиксируются один раз. Исходный датасет в архив не включён: для воспроизведения используется файл организаторов.')
    p('Запуск: python -m venv .venv; установка requirements.txt; python src/finalize.py --rebuild; python -m unittest discover -s tests -v; python src/validate_artifacts.py. Демо: streamlit run app.py. Точные команды и структура приведены в README.')
    p('Сдача — 29 октября 2026 года, до 21:00 по Москве. Промежуточный официальный скоринг не предоставляется.')
    SimpleDocTemplate(str(OUT/'report.pdf'),pagesize=A4,rightMargin=42,leftMargin=42,topMargin=35,bottomMargin=35).build(story)


def build_slides(meta):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    blue=RGBColor(23,60,102);gray=RGBColor(68,83,104)
    text=narrative(meta)
    def box(slide,x,y,w,h,content,size=24,bold=False,color=gray):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=shape.text_frame;tf.word_wrap=True
        for i,line in enumerate(content.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
            p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color;p.space_after=Pt(12)
        return shape
    def slide(title,number):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        box(s,.6,.35,12.1,.8,title,30,True,blue)
        box(s,.6,7.03,12,.25,f'РОСХИМ · Прогноз цены ПВХ · 29 октября 2026     |     {number}/5',10,color=gray)
        return s
    s=slide('ПВХ: прогноз цены и поддержка решений',1)
    box(s,.7,1.5,11.8,1.4,'SG-5 (Китай), DAP Достык\n13 будущих недель · анонимизированный индекс k',30,True,blue)
    box(s,.7,3.3,11.7,2.7,'Прогноз + интервалы + фиксированная стратегия + интерактивное демо\nДля обсуждения сроков продаж, закупок и планирования\nРабочее воспроизводимое решение; гарантированная прибыль не заявляется',24)
    s=slide('Метод: только информация из прошлого',2)
    box(s,.7,1.5,11.9,1.,'XLSB → календарь → признаки → временная проверка → прогноз',27,True,blue)
    box(s,.7,2.8,11.8,3.7,'36 293 записи · 170 рядов · недельная цель\n24 блока по 13 недель; признаки и обучение внутри каждого префикса\n16 исходных вариантов: naive, Holt, Ridge, CatBoost и ансамбли\nФинальный кандидат: CatBoost MAE с рыночными признаками\nДополнительно проверены 2 осторожных ансамбля с naive',23)
    s=slide('Качество: небольшой выигрыш, устойчивость ограничена',3)
    short=(f"Последние 8 окон: {meta['final_model_recent8_smape']:.2f}% против {meta['final_naive_recent8_smape']:.2f}% у naive.\n"
           f"20 блоков: {meta['final_model_mean_smape']:.2f}% против {meta['final_naive_mean_smape']:.2f}% у naive.\n"
           'Устойчивое превосходство не подтверждено.')
    box(s,.7,1.35,11.8,1.65,short,20)
    names=['Последние 8 окон','20 исторических блоков']
    fig,ax=plt.subplots(figsize=(9,2.5));x=[0,1]
    ax.bar([i-.18 for i in x],[meta['final_model_recent8_smape'],meta['final_model_mean_smape']],width=.35,label='Frozen CatBoost')
    ax.bar([i+.18 for i in x],[meta['final_naive_recent8_smape'],meta['final_naive_mean_smape']],width=.35,label='Naive',color='#94a3b8')
    ax.set_xticks(x,['Recent 8 blocks (selection)','20 blocks (development)']);ax.set_ylabel('sMAPE, %');ax.legend();fig.tight_layout()
    chart=OUT/'model_comparison.png';fig.savefig(chart,dpi=160);plt.close(fig)
    s.shapes.add_picture(str(chart),Inches(1.3),Inches(3.2),width=Inches(10.6),height=Inches(2.95))
    box(s,.7,6.2,11.8,.55,'Ансамбли снижали общий риск ошибки, но ухудшали недавние окна: прогноз оставлен без замены.',17)
    s=slide('Ошибка растёт с горизонтом — пересмотр полезен',4)
    table=pd.read_csv(ART/'balanced_horizon_metrics.csv')
    fig,ax=plt.subplots(figsize=(9,3.1))
    for name,label in [(meta['model'],'Frozen CatBoost'),('naive','Naive')]:
        g=table[table.model.eq(name)];ax.plot(g.horizon,g.smape,marker='o',label=label)
    ax.set(xlabel='Horizon, weeks',ylabel='sMAPE, %',xticks=list(range(1,14)));ax.grid(alpha=.25);ax.legend();fig.tight_layout()
    chart=OUT/'horizon.png';fig.savefig(chart,dpi=160);plt.close(fig)
    s.shapes.add_picture(str(chart),Inches(1.0),Inches(1.45),width=Inches(11.2),height=Inches(3.8))
    box(s,.7,5.35,11.8,1.2,'Уточнение t+2 после t+1: 4,44% → 2,87% sMAPE на 15 парах.\nЭто отдельное исследование: конкурсный прогноз новых фактов не получает.',19)
    s=slide('Риск и готовый инструмент',5)
    box(s,.7,1.45,11.8,3.7,f"Интервалы: покрытие {meta['interval_coverage']:.1%}, широкие, без гарантии.\nСтратегия: {meta['strategy_pnl']:+.3f}; после условных затрат {meta['strategy_stress_pnl']:+.3f}.\nМаксимум одна активная неделя на 13 недель.\n{meta.get('strategy_profitable_blocks',0)} прибыльных блоков; лучший — {meta.get('strategy_largest_positive_fraction',0):.0%} положительного результата.\nЭто историческая оценка; будущая прибыль не гарантируется.",23)
    box(s,.7,5.6,11.8,1.,'Следующий шаг после конкурса: новые котировки, устойчивый выбор режима и проверка экономической пользы.',22,True,blue)
    prs.save(OUT/'presentation.pptx')


def package(meta):
    readme='''КОМПЛЕКТ СДАЧИ — РОСХИМ

forecast.csv — официальный прогноз с границами интервала.
strategy.csv — фиксированные позиции: -1 в первую неделю, затем 0.
report.pdf — отчёт на двух страницах.
presentation.pptx — пять слайдов для защиты; проверьте оформление на своей машине.
code/ — исходный код, инструкции, шаблоны и параметры воспроизведения.
model/model.cbm — финальная модель CatBoost (прогнозирует логарифм изменения k).
validation/ — сводные метрики и протокол финального выбора без исходных цен.

Файлы CSV: ;, UTF-8 BOM, 13 пятниц 10.07–02.10.2026.
Активная стратегия ограничена одной неделей; исторический плюс не гарантирует будущую прибыль.
Датасет не включён. Для воспроизведения поместите официальный файл
в code/data/raw/predictors.xlsb и следуйте code/README.md.
Канонический расчёт: python src/finalize.py --rebuild.
Сдача: 29 октября до 21:00 по Москве.
'''
    (OUT/'README_DELIVERY.txt').write_text(readme)
    files={name:ART/name for name in ['forecast.csv','strategy.csv']}
    files.update({name:OUT/name for name in ['report.pdf','presentation.pptx','README_DELIVERY.txt']})
    model=ART/'model/model.cbm'
    if model.exists():files['model/model.cbm']=model
    for directory in ['src','tests','config','templates','docs']:
        for p in (ROOT/directory).rglob('*'):
            if p.is_file() and '__pycache__' not in p.parts and p.suffix not in ['.pdf','.pptx']:
                files['code/'+str(p.relative_to(ROOT))]=p
    for name in ['README.md','requirements.txt','app.py','.python-version','.gitignore','data/raw/.gitkeep']:
        files['code/'+name]=ROOT/name
    for name in ['run_metadata.json','final_model_comparison.csv','final_strategy_comparison.csv','final_strategy_folds.csv','interval_metrics.csv','revision_metrics.csv']:
        files['validation/'+name]=ART/name
    manifest={name:hashlib.sha256(p.read_bytes()).hexdigest() for name,p in sorted(files.items())}
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    files['manifest.json']=OUT/'manifest.json'
    with zipfile.ZipFile(OUT/'Roshim_submission.zip','w',zipfile.ZIP_DEFLATED) as z:
        for name,path in sorted(files.items()):z.write(path,arcname=name)
    print('Built',OUT/'Roshim_submission.zip',len(files),'files')


if __name__=='__main__':
    OUT.mkdir(exist_ok=True)
    meta=json.loads((ART/'run_metadata.json').read_text())
    if not meta.get('finalized'):
        raise SystemExit('Run src/finalize.py before packaging')
    build_pdf(meta);build_slides(meta)
    shutil.copy2(OUT/'report.pdf',ROOT/'docs/report.pdf')
    shutil.copy2(OUT/'presentation.pptx',ROOT/'docs/presentation.pptx')
    package(meta)
